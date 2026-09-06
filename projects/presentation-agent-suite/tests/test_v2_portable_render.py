"""Actual portable PPTX rendering, with explicit absence/cancellation boundaries."""
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
from presentation_agents import pipeline as legacy
from presentation_agents.v2 import portable_render


class PortableRenderTests(unittest.TestCase):
    def setUp(self):
        runtime=Path(__file__).resolve().parents[3]/'.runtime';runtime.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=runtime,prefix='portable-render-test-')
        self.root = Path(self.temp.name)
        self.pptx = self.root/'source.pptx'
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        for number in range(1,4):
            slide=presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(2)).text=f'Portable renderer page {number}'
        presentation.save(self.pptx)
        self.contact=self.root/'render/grid.png'

    def tearDown(self): self.temp.cleanup()

    @unittest.skipUnless(portable_render.find_soffice(), 'Actual bundled/installed LibreOffice required')
    def test_actual_pptx_has_three_ordered_high_resolution_pages_and_matching_hashes(self):
        before=self.pptx.read_bytes()
        proof=portable_render.render_pptx(self.pptx,self.contact)
        self.assertEqual(proof['source_sha256'],hashlib.sha256(before).hexdigest())
        self.assertEqual(proof['slide_count'],3)
        self.assertEqual([p['page'] for p in proof['pages']],[1,2,3])
        hashes=[]
        for page in proof['pages']:
            data=(self.contact.parent/page['image_path']).read_bytes()
            self.assertGreaterEqual(page['image_width'],1920)
            self.assertEqual(page['image_sha256'],hashlib.sha256(data).hexdigest())
            hashes.append(page['image_sha256'])
        self.assertEqual(len(set(hashes)),3)
        self.assertEqual(proof,json.loads(self.contact.with_suffix('.render.json').read_text()))
        self.assertEqual(before,self.pptx.read_bytes())

    def test_success_exit_without_pdf_is_not_a_render_and_leaves_no_pages(self):
        with patch.dict(portable_render.os.environ, {'PYTHONDONTWRITEBYTECODE': '0'}), patch.object(portable_render,'find_soffice',return_value='/fixture/soffice'), patch.object(legacy,'_run_bounded_subprocess',return_value=subprocess.CompletedProcess([],0)) as render:
            with self.assertRaisesRegex(RuntimeError,'bounded PDF'):
                portable_render.render_pptx(self.pptx,self.contact)
            self.assertEqual(render.call_args.kwargs['env']['PYTHONDONTWRITEBYTECODE'], '1')
        self.assertFalse(self.contact.exists())
        self.assertFalse((self.contact.parent/'grid-pages').exists())

    def test_changed_copy_cannot_receive_original_source_provenance_after_restore(self):
        import pymupdf
        original=self.pptx.read_bytes()
        # Simulates source A -> B during copying -> A before the final read.
        # No real LibreOffice or model is invoked in this adversarial fixture.
        def changed_copy(source,destination):
            Path(destination).write_bytes(b'candidate B fixture')
            self.pptx.write_bytes(original)
        def render_other_candidate(command,**kwargs):
            pdf=Path(kwargs['cwd'])/'candidate.pdf'
            with pymupdf.open() as document:
                document.new_page(width=960,height=540).insert_text((30,30),'Candidate B')
                document.save(pdf)
            return subprocess.CompletedProcess(command,0)
        with patch.object(portable_render,'find_soffice',return_value='/fixture/soffice'),patch.object(portable_render.shutil,'copyfile',side_effect=changed_copy),patch.object(legacy,'_run_bounded_subprocess',side_effect=render_other_candidate) as render:
            with self.assertRaisesRegex(ValueError,'snapshot|copy'):
                portable_render.render_pptx(self.pptx,self.contact)
            render.assert_not_called()
        self.assertFalse(self.contact.exists());self.assertFalse((self.contact.parent/'grid-pages').exists())
        self.assertEqual(self.pptx.read_bytes(),original)

    def test_cancellation_before_child_launch_does_not_leave_reusable_outputs(self):
        with patch.object(portable_render,'find_soffice',return_value='/fixture/soffice'):
            with self.assertRaises(legacy.SubprocessCancelled):
                portable_render.render_pptx(self.pptx,self.contact,cancelled=lambda:True)
        self.assertFalse(self.contact.exists())
        self.assertFalse((self.contact.parent/'grid-pages').exists())


class KoreanFontRenderTests(unittest.TestCase):
    @unittest.skipUnless(portable_render.find_soffice(), 'Actual bundled LibreOffice required')
    def test_actual_final_svg_candidate_renders_korean_glyphs_without_font_install(self):
        import sys
        from PIL import Image
        repo = Path(__file__).resolve().parents[3]
        scripts = repo / '.claude/skills/ppt-master/scripts'
        if str(scripts) not in sys.path: sys.path.insert(0, str(scripts))
        from svg_to_pptx.pptx_package.builder import create_pptx_with_native_svg
        runtime = repo / '.runtime'; runtime.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='korean-native-render-', dir=runtime) as temporary:
            root = Path(temporary)
            svg = root / 'P01.svg'
            # Korean only: an empty/tofu-free text extraction cannot pass by
            # drawing the Latin label while silently omitting the Hangul.
            svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><rect width="1280" height="720" fill="#FFFFFF"/><text x="80" y="180" font-family="Pretendard" font-size="48" fill="#111111">한글은퇴생계설계가나다힣</text></svg>', encoding='utf-8')
            candidate = root / 'candidate.pptx'
            self.assertTrue(create_pptx_with_native_svg([svg], candidate, pptx_structure='flat', verbose=False, transition=None, enable_notes=False))
            original = candidate.read_bytes()
            contact = root / 'render/grid.png'
            actual_run = legacy._run_bounded_subprocess
            used_fonts = set()
            def observe_pdf(command, **kwargs):
                import pymupdf
                result = actual_run(command, **kwargs)
                pdf = Path(kwargs['cwd']) / 'candidate.pdf'
                if pdf.is_file():
                    with pymupdf.open(pdf) as document:
                        used_fonts.update(font[3].split('+')[-1] for font in document[0].get_fonts())
                return result
            with patch.object(legacy, '_run_bounded_subprocess', side_effect=observe_pdf):
                proof = portable_render.render_pptx(candidate, contact)
            self.assertIn('Pretendard-Regular', used_fonts)
            self.assertEqual(proof['source_sha256'], hashlib.sha256(original).hexdigest())
            self.assertEqual(candidate.read_bytes(), original)
            self.assertEqual(proof['slide_count'], 1)
            with Image.open(contact.parent / proof['pages'][0]['image_path']) as page:
                # Every syllable must occupy visible ink in the actual PNG.
                # Measured independently from the pinned original OTF at 48px:
                # Pretendard Hangul advances are 41.484375 SVG pixels.
                scale = page.width / 1280
                advance = 41.484375
                for index in range(len('한글은퇴생계설계가나다힣')):
                    left = int((80 + index * advance) * scale)
                    box = (left, int(128 * scale), int((80 + (index + 1) * advance) * scale), int(184 * scale))
                    dark = sum(count for value, count in enumerate(page.crop(box).convert('L').histogram()) if value < 100)
                    self.assertGreater(dark, 60, f'Hangul glyph {index + 1} is absent in the native PNG')


if __name__=='__main__': unittest.main()
