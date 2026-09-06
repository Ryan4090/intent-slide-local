"""Real package isolation/composition; native renderer processes are simulated here."""
import io
import json
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from lxml import etree
from PIL import Image
from pptx import Presentation

from presentation_agents import pipeline as legacy


class QuickLookIsolationTests(unittest.TestCase):
    def setUp(self):
        runtime = Path(__file__).resolve().parents[1] / '.runtime'
        runtime.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=runtime, prefix='quicklook-test-')
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.pptx'
        deck = Presentation()
        deck.slide_width, deck.slide_height = 1280 * 9525, 720 * 9525
        for index in range(3):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = f'Page {index + 1}'
            slide.notes_slide.notes_text_frame.text = f'Notes {index + 1}'
        deck.save(self.source)
        self.before = self.source.read_bytes()
        self.output = self.root / 'contact.png'

    def tearDown(self):
        self.temp.cleanup()

    def test_isolation_changes_only_presentation_slide_selection(self):
        isolated = self.root / 'page2.pptx'
        descriptor = legacy._write_quicklook_single_slide(self.source, isolated, 1)
        with zipfile.ZipFile(self.source) as old, zipfile.ZipFile(isolated) as new:
            self.assertEqual(old.namelist(), new.namelist())
            changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
            self.assertEqual(changed, ['ppt/presentation.xml'])
            ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
            old_slides = etree.fromstring(old.read('ppt/presentation.xml')).find('p:sldIdLst', ns)
            new_slides = etree.fromstring(new.read('ppt/presentation.xml')).find('p:sldIdLst', ns)
            self.assertEqual(len(new_slides), 1)
            self.assertEqual(dict(new_slides[0].attrib), dict(old_slides[1].attrib))
        self.assertEqual(descriptor['page'], 2)
        self.assertEqual(self.source.read_bytes(), self.before)

    @staticmethod
    def _write_render_fixture(command, color='white'):
        Image.new('RGB', (504, 320), color).save(command[3])
        directory = Path(command[5])
        directory.mkdir()
        Image.new('RGB', (1920, 1080), color).save(directory / 'P01.png')

    def _run(self, external, cancelled=None):
        with patch.object(legacy.sys, 'platform', 'darwin'), patch.object(
            legacy.shutil, 'which', side_effect=lambda name: '/fixture/' + name
        ), patch.object(legacy, '_run_bounded_subprocess', side_effect=external):
            return legacy._render_macos_quicklook_contact_sheet(self.source, self.output, cancelled=cancelled)

    def test_every_page_isolated_in_order_with_bound_provenance(self):
        calls = []
        colors = [(200, 20, 20), (20, 180, 20), (20, 20, 200)]
        def external(command, **kwargs):
            calls.append(command)
            if 'qlmanage' in command[0]:
                with zipfile.ZipFile(command[-1]) as archive:
                    tree = etree.fromstring(archive.read('ppt/presentation.xml'))
                    self.assertEqual(len(tree.find('{*}sldIdLst')), 1)
                (Path(command[command.index('-o') + 1]) / 'Preview.html').write_text('fixture')
            else:
                self._write_render_fixture(command, colors[len(calls) // 2 - 1])
            return subprocess.CompletedProcess(command, 0, '', '')
        self.assertTrue(self._run(external))
        self.assertEqual(len(calls), 6)
        proof = json.loads(self.output.with_suffix('.render.json').read_text())
        self.assertEqual(proof['method'], 'isolated-slide-selection')
        self.assertEqual([p['page'] for p in proof['pages']], [1, 2, 3])
        self.assertEqual(proof['source_sha256'], legacy.sha256_file(self.source))
        self.assertEqual(proof['contact_sheet_sha256'], legacy.sha256_file(self.output))
        with Image.open(self.output) as grid:
            for index, color in enumerate(colors):
                self.assertEqual(grid.getpixel((12 + index * 492 + 100, 100)), color)
        self.assertEqual(self.source.read_bytes(), self.before)
        self.assertFalse(list(self.root.glob('presentation-quicklook-*')))

    def test_failure_does_not_publish_partial_grid_or_provenance(self):
        calls = 0
        def external(command, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 3:
                return subprocess.CompletedProcess(command, 1, '', 'fixture failure')
            if 'qlmanage' in command[0]:
                (Path(command[command.index('-o') + 1]) / 'Preview.html').write_text('fixture')
            else:
                self._write_render_fixture(command)
            return subprocess.CompletedProcess(command, 0, '', '')
        self.assertFalse(self._run(external))
        self.assertFalse(self.output.exists())
        self.assertFalse(self.output.with_suffix('.render.json').exists())
        self.assertFalse((self.root / 'contact-pages').exists())

    def test_all_pages_use_one_snapshot_if_source_changes_then_is_restored(self):
        changed = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(self.before)) as original, zipfile.ZipFile(changed, 'w') as alternate:
            for entry in original.infolist():
                data = original.read(entry.filename)
                if entry.filename == 'ppt/slides/slide2.xml':
                    self.assertIn(b'Page 2', data)
                    data = data.replace(b'Page 2', b'CHANGED PAGE 2')
                alternate.writestr(entry, data)
        calls = 0
        selected_second_page = None

        def external(command, **kwargs):
            nonlocal calls, selected_second_page
            calls += 1
            if 'qlmanage' in command[0]:
                if calls == 3:
                    with zipfile.ZipFile(command[-1]) as selected:
                        selected_second_page = selected.read('ppt/slides/slide2.xml')
                (Path(command[command.index('-o') + 1]) / 'Preview.html').write_text('fixture')
            else:
                self._write_render_fixture(command)
                if calls == 2:
                    self.source.write_bytes(changed.getvalue())
                elif calls == 4:
                    self.source.write_bytes(self.before)
            return subprocess.CompletedProcess(command, 0, '', '')

        self.assertTrue(self._run(external))
        self.assertIsNotNone(selected_second_page)
        self.assertIn(b'Page 2', selected_second_page)
        self.assertNotIn(b'CHANGED PAGE 2', selected_second_page)
        self.assertEqual(self.source.read_bytes(), self.before)
        proof = json.loads(self.output.with_suffix('.render.json').read_text())
        self.assertEqual(proof['source_sha256'], legacy.sha256_file(self.source))
        self.assertFalse(list(self.root.glob('presentation-quicklook-*')))

    def test_cancel_between_pages_removes_temporary_outputs(self):
        state = {'cancel': False}
        callback = lambda: state['cancel']
        def external(command, **kwargs):
            self.assertIs(kwargs['cancelled'], callback)
            if 'qlmanage' in command[0]:
                (Path(command[command.index('-o') + 1]) / 'Preview.html').write_text('fixture')
            else:
                self._write_render_fixture(command)
                state['cancel'] = True
            return subprocess.CompletedProcess(command, 0, '', '')
        with self.assertRaises(legacy.SubprocessCancelled):
            self._run(external, callback)
        self.assertFalse(self.output.exists())
        self.assertFalse(self.output.with_suffix('.render.json').exists())
        self.assertFalse(self.output.with_name('contact-pages').exists())
        self.assertFalse(list(self.root.glob('presentation-quicklook-*')))

    def test_invalid_page_selection_is_rejected(self):
        for index in (-1, 3):
            with self.subTest(index=index), self.assertRaises(legacy.ContractError):
                legacy._write_quicklook_single_slide(self.source, self.root / 'bad.pptx', index)


if __name__ == '__main__':
    unittest.main()
