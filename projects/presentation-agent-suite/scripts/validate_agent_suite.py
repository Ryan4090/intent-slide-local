#!/usr/bin/env python3
"""
Presentation Agent Suite - Contract Validator

Exercises the three-agent handoff without contacting an AI provider or the web.

Usage:
    .venv/bin/python projects/presentation-agent-suite/scripts/validate_agent_suite.py

Dependencies:
    PyMuPDF from the repository-local virtual environment.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
import zlib
from pathlib import Path
from unittest.mock import patch

_SUITE_ROOT = Path(__file__).resolve().parents[1]
_REPO_ROOT = _SUITE_ROOT.parents[1]
_SRC_DIR = _SUITE_ROOT / "src"
_AGENT_LAUNCHER = _SUITE_ROOT / "scripts/presentation_agents.py"
_SUPERVISOR_LAUNCHER = _SUITE_ROOT / "scripts/presentation_supervisor.py"
if str(_SRC_DIR) not in sys.path:
    sys.path.insert(0, str(_SRC_DIR))

import presentation_agents.pipeline as pipeline_module  # noqa: E402
from presentation_agents.pipeline import (  # noqa: E402
    AgentPipeline,
    ContractError,
    _validate_png_bytes,
)
from presentation_agents.pdf_report import render_markdown_pdf  # noqa: E402
from presentation_agents.utils import sha256_file  # noqa: E402
from presentation_agents.validation import validate_agent_manifests  # noqa: E402


def _confirmed(value: object, source: str = "user") -> dict[str, object]:
    return {
        "value": value,
        "state": "confirmed",
        "source": source,
    }


def _sample_intent() -> dict[str, object]:
    return {
        "schema_version": "1.0.0",
        "fields": {
            "topic": _confirmed("2027 체외진단 시장 진입 전략"),
            "audience": _confirmed("아시아태평양 사업부 경영진"),
            "objective": _confirmed("우선 진입 시장 세 곳의 투자 타당성을 검토한다"),
            "desired_decision": _confirmed("1순위 시장과 검증 예산을 승인한다"),
            "key_message": _confirmed("시장 매력도보다 실행 가능성과 증거 품질을 함께 봐야 한다"),
            "call_to_action": _confirmed("1순위 시장의 90일 검증 계획을 승인한다"),
            "source_scope": _confirmed("공개 기관·규제기관·기업 공시와 사용자가 제공한 내부 자료"),
            "must_include": _confirmed(["시장 규모", "규제", "경쟁", "90일 계획"]),
            "must_exclude": _confirmed(["출처 없는 시장 전망", "확정되지 않은 내부 수치"]),
            "delivery_mode": _confirmed("balanced"),
            "duration_minutes": _confirmed(15),
            "slide_count": _confirmed(5),
            "language": _confirmed("ko-KR"),
            "content_divergence": _confirmed("근거는 유지하되 경영진 의사결정 흐름으로 재구성"),
            "template_or_brand": _confirmed("free design; Pretendard"),
            "speaker_notes": _confirmed(False),
            "confidentiality": _confirmed("공개 자료와 합성 검증 데이터만 사용"),
            "success_criteria": _confirmed(["시장 우선순위가 근거로 추적됨", "90일 실행 결정 가능"]),
        },
    }


def _sample_outline() -> dict[str, object]:
    return {
        "schema_version": "1.0.0",
        "status": "approved_candidate",
        "slides": [
            {
                "slide_id": "P01",
                "role": "context",
                "title": "어디에 먼저 진입할 것인가",
                "purpose": "의사결정 질문과 평가 기준을 고정한다",
                "content": ["후보 시장 세 곳", "매력도·규제·실행 가능성의 통합 평가"],
                "evidence_needed": ["후보 시장 정의"],
                "visual_intent": "의사결정 프레임",
            },
            {
                "slide_id": "P02",
                "role": "evidence",
                "title": "수요는 성장하지만 시장별 증거의 질은 다르다",
                "purpose": "시장 기회와 불확실성을 분리한다",
                "content": ["시장별 규모와 성장", "데이터 한계"],
                "evidence_needed": ["시장 규모", "전망 범위와 기준연도"],
                "visual_intent": "small multiples",
            },
            {
                "slide_id": "P03",
                "role": "comparison",
                "title": "규제와 채널이 실제 진입 속도를 가른다",
                "purpose": "실행 장벽을 비교한다",
                "content": ["승인 경로", "유통·상환 경로", "경쟁 밀도"],
                "evidence_needed": ["규제기관 원문", "채널 근거"],
                "visual_intent": "evidence matrix",
            },
            {
                "slide_id": "P04",
                "role": "finding",
                "title": "1순위 후보는 시장 A다",
                "purpose": "평가 결과와 민감도를 제시한다",
                "content": ["가중 점수", "가정 변화 시 순위"],
                "evidence_needed": ["평가표", "가중치 근거"],
                "visual_intent": "ranked scorecard",
            },
            {
                "slide_id": "P05",
                "role": "decision",
                "title": "90일 검증으로 투자 결정을 좁힌다",
                "purpose": "승인할 행동과 종료 조건을 명시한다",
                "content": ["3개 workstream", "결정 게이트", "중단 조건"],
                "evidence_needed": ["실행 제약", "성공 기준"],
                "visual_intent": "roadmap",
            },
        ],
    }


def _expect_blocked(action, contains: str) -> None:
    try:
        action()
    except ContractError as exc:
        if contains not in str(exc):
            raise AssertionError(f"expected {contains!r} in {exc!r}") from exc
    else:
        raise AssertionError(f"expected ContractError containing {contains!r}")


def _png_chunk(chunk_type: bytes, payload: bytes) -> bytes:
    crc = zlib.crc32(chunk_type)
    crc = zlib.crc32(payload, crc) & 0xFFFFFFFF
    return struct.pack(">I", len(payload)) + chunk_type + payload + struct.pack(">I", crc)


def _valid_png() -> bytes:
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    scanline = zlib.compress(b"\x00\x00\x00\x00\x00")
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", ihdr)
        + _png_chunk(b"IDAT", scanline)
        + _png_chunk(b"IEND", b"")
    )


def _illegal_truecolor_png() -> bytes:
    ihdr = struct.pack(">IIBBBBB", 1, 1, 1, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", ihdr)
        + _png_chunk(b"IDAT", zlib.compress(b"\x00\x00"))
        + _png_chunk(b"IEND", b"")
    )


def _indexed_png_without_palette() -> bytes:
    ihdr = struct.pack(">IIBBBBB", 1, 1, 1, 3, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", ihdr)
        + _png_chunk(b"IDAT", zlib.compress(b"\x00\x00"))
        + _png_chunk(b"IEND", b"")
    )


def _nonconsecutive_idat_png() -> bytes:
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    compressed = zlib.compress(b"\x00\x00\x00\x00\x00")
    split_at = len(compressed) // 2
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", ihdr)
        + _png_chunk(b"IDAT", compressed[:split_at])
        + _png_chunk(b"tEXt", b"probe")
        + _png_chunk(b"IDAT", compressed[split_at:])
        + _png_chunk(b"IEND", b"")
    )


def _verify_quicklook_helper_contract(root: Path) -> None:
    if sys.platform != "darwin":
        return
    renderer = _SUITE_ROOT / "scripts/render_quicklook_contact_sheet.swift"
    if not renderer.is_file():
        return
    from PIL import Image
    from pptx import Presentation

    pptx = root / "helper.pptx"
    output = root / "helper-grid.png"
    provenance = output.with_suffix(".render.json")
    fixture = Presentation()
    fixture.slide_width, fixture.slide_height = 1280 * 9525, 720 * 9525
    fixture.slides.add_slide(fixture.slide_layouts[6])
    fixture.save(pptx)
    original_bytes = pptx.read_bytes()
    tool_paths = {"qlmanage": "/usr/bin/qlmanage", "swift": "/usr/bin/swift"}
    calls: list[tuple[list[str], dict[str, object]]] = []
    preview_parent: Path | None = None
    preview_html: Path | None = None

    def _which(name: str) -> str | None:
        return tool_paths.get(name)

    def _successful_run(
        command: list[str], **kwargs: object
    ) -> subprocess.CompletedProcess[str]:
        nonlocal preview_parent, preview_html
        calls.append((command, kwargs))
        if command[0] == tool_paths["qlmanage"]:
            assert command[1:3] == ["-p", "-o"]
            preview_root = Path(command[3])
            preview_parent = preview_root.parent
            selected = Path(command[4])
            assert selected != pptx and selected.parent == preview_parent
            assert len(Presentation(selected).slides) == 1
            package = preview_root / "helper.pptx.qlpreview"
            package.mkdir(parents=True)
            preview_html = package / "Preview.html"
            preview_html.write_text("<html></html>", encoding="utf-8")
        else:
            assert preview_html is not None
            assert command == [
                tool_paths["swift"],
                str(renderer),
                str(preview_html),
                str(preview_parent / "page.png"),
                "--slide-images-dir", str(preview_parent / "slide-images"),
            ]
            Image.new("RGB", (504, 320), "white").save(command[3])
            Path(command[5]).mkdir()
            Image.new("RGB", (1920, 1080), "white").save(Path(command[5]) / "P01.png")
        return subprocess.CompletedProcess(command, 0, stdout="PASS", stderr="")

    with patch("presentation_agents.pipeline.shutil.which", side_effect=_which), patch(
        "presentation_agents.pipeline._run_bounded_subprocess", side_effect=_successful_run
    ), patch(
        "presentation_agents.pipeline.subprocess.run",
        side_effect=AssertionError("Quick Look must not use unbounded subprocess.run"),
    ):
        assert pipeline_module._render_macos_quicklook_contact_sheet(pptx, output)
    assert len(calls) == 2
    assert 0 < calls[1][1]["timeout"] <= calls[0][1]["timeout"] <= 600
    assert preview_parent is not None and not preview_parent.exists()
    assert pptx.read_bytes() == original_bytes
    proof = json.loads(provenance.read_text(encoding="utf-8"))
    assert proof["source_sha256"] == sha256_file(pptx)
    assert proof["contact_sheet_sha256"] == sha256_file(output)
    assert proof["slide_count"] == 1 and proof["pages"][0]["page"] == 1
    output.unlink()
    provenance.unlink()

    def _quicklook_with_preview_count(
        count: int,
        *,
        second_returncode: int = 0,
        invalid_png: bool = False,
        timeout_second: bool = False,
    ):
        invocation = 0

        def _run(command: list[str], **_: object) -> subprocess.CompletedProcess[str]:
            nonlocal invocation
            invocation += 1
            if invocation == 1:
                preview_root = Path(command[3])
                for index in range(count):
                    package = preview_root / f"package-{index}.qlpreview"
                    package.mkdir(parents=True)
                    (package / "Preview.html").write_text("<html></html>", encoding="utf-8")
                return subprocess.CompletedProcess(command, 0, stdout="PASS", stderr="")
            if timeout_second:
                raise subprocess.TimeoutExpired(command, 600)
            if invalid_png:
                Path(command[3]).write_bytes(b"invalid")
            else:
                Image.new("RGB", (504, 320), "white").save(command[3])
                Path(command[5]).mkdir()
                Image.new("RGB", (1920, 1080), "white").save(Path(command[5]) / "P01.png")
            return subprocess.CompletedProcess(
                command, second_returncode, stdout="", stderr="synthetic renderer failure"
            )

        return _run

    for count in (0, 2):
        with patch("presentation_agents.pipeline.shutil.which", side_effect=_which), patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_quicklook_with_preview_count(count),
        ):
            assert not pipeline_module._render_macos_quicklook_contact_sheet(pptx, output)
        assert not output.exists()
        assert not provenance.exists()

    with patch("presentation_agents.pipeline.shutil.which", side_effect=_which), patch(
        "presentation_agents.pipeline._run_bounded_subprocess",
        side_effect=subprocess.TimeoutExpired(["qlmanage"], 600),
    ):
        assert not pipeline_module._render_macos_quicklook_contact_sheet(pptx, output)
    assert not output.exists()
    assert not provenance.exists()

    for run in (
        _quicklook_with_preview_count(1, timeout_second=True),
        _quicklook_with_preview_count(1, second_returncode=1),
    ):
        with patch("presentation_agents.pipeline.shutil.which", side_effect=_which), patch(
            "presentation_agents.pipeline._run_bounded_subprocess", side_effect=run
        ):
            assert not pipeline_module._render_macos_quicklook_contact_sheet(pptx, output)
        assert not output.exists()
        assert not provenance.exists()

    diagnostics: list[str] = []
    with patch("presentation_agents.pipeline.shutil.which", side_effect=_which), patch(
        "presentation_agents.pipeline._run_bounded_subprocess",
        side_effect=_quicklook_with_preview_count(1, invalid_png=True),
    ):
        assert not pipeline_module._render_macos_quicklook_contact_sheet(
            pptx,
            output,
            diagnostics=diagnostics,
        )
    assert any("invalid contact-sheet PNG" in item for item in diagnostics)
    assert not output.exists()
    assert not provenance.exists()
    assert pptx.read_bytes() == original_bytes


def _verify_bounded_subprocess_contract() -> None:
    command = [
        sys.executable,
        "-c",
        (
            "import sys; "
            "sys.stdout.write('x' * 200000); "
            "sys.stderr.write('y' * 200000); "
            "raise SystemExit(7)"
        ),
    ]
    result = pipeline_module._run_bounded_subprocess(command, timeout=30)
    assert result.returncode == 7
    assert len(result.stdout.encode("utf-8")) <= 66_000
    assert len(result.stderr.encode("utf-8")) <= 66_000
    assert "output truncated" in result.stdout
    assert "output truncated" in result.stderr

    verify_deck = _load_verify_deck_module()
    owner_result = verify_deck._run_bounded_subprocess(command, timeout=30)
    assert owner_result.returncode == 7
    assert len(owner_result.stdout.encode("utf-8")) <= 66_000
    assert len(owner_result.stderr.encode("utf-8")) <= 66_000
    assert "output truncated" in owner_result.stdout
    assert "output truncated" in owner_result.stderr

    timeout_command = [
        sys.executable,
        "-c",
        (
            "import sys,time; "
            "sys.stdout.write('x' * 200000); sys.stdout.flush(); "
            "sys.stderr.write('y' * 200000); sys.stderr.flush(); "
            "time.sleep(30)"
        ),
    ]
    for runner in (
        pipeline_module._run_bounded_subprocess,
        verify_deck._run_bounded_subprocess,
    ):
        try:
            runner(timeout_command, timeout=1)
        except subprocess.TimeoutExpired as exc:
            stdout = str(exc.output or "")
            stderr = str(exc.stderr or "")
            assert len(stdout.encode("utf-8")) <= 66_000
            assert len(stderr.encode("utf-8")) <= 66_000
            assert "output truncated" in stdout
            assert "output truncated" in stderr
        else:
            raise AssertionError("bounded subprocess must enforce timeout")


def _load_verify_deck_module():
    script = _REPO_ROOT / ".claude/skills/ppt-master/scripts/verify_deck.py"
    spec = importlib.util.spec_from_file_location("suite_verify_deck_contract", script)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _verify_officecli_render_failure_contract(root: Path) -> None:
    verify_deck = _load_verify_deck_module()
    project = root / "officecli-render-contract"
    exports = project / "exports"
    exports.mkdir(parents=True)
    pptx = exports / "fixture.pptx"
    pptx.write_bytes(b"fixture")
    output = project / "_pptx_render/fixture-grid.png"

    def _failed_after_output(args: list[str], timeout: int):
        assert timeout in {120, 180}
        if "validate" in args:
            return subprocess.CompletedProcess(
                args,
                0,
                stdout=json.dumps({"success": True}),
                stderr="",
            )
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(_valid_png())
        return subprocess.CompletedProcess(
            args,
            9,
            stdout="",
            stderr="synthetic OfficeCLI render failure",
        )

    with patch.object(verify_deck, "_officecli_bin", return_value="/usr/bin/officecli"), patch.object(
        verify_deck,
        "_run_officecli",
        side_effect=_failed_after_output,
    ):
        failures, warnings = verify_deck.officecli_checks(
            project,
            render=True,
            require_render_success=True,
        )
    assert any("officecli screenshot failed" in item for item in failures), failures
    assert not any("officecli screenshot failed" in item for item in warnings), warnings
    assert not output.exists()

    invocations = 0

    def _timeout(args: list[str], timeout: int):
        nonlocal invocations
        invocations += 1
        if invocations == 1:
            return subprocess.CompletedProcess(
                args,
                0,
                stdout=json.dumps({"success": True}),
                stderr="",
            )
        return None

    with patch.object(verify_deck, "_officecli_bin", return_value="/usr/bin/officecli"), patch.object(
        verify_deck,
        "_run_officecli",
        side_effect=_timeout,
    ):
        failures, _ = verify_deck.officecli_checks(
            project,
            render=True,
            require_render_success=True,
        )
    assert any("officecli screenshot failed" in item for item in failures), failures


def _verify_hash_receipt_regressions() -> list[str]:
    """Prove cache invalidation without modifying shared validator sources."""
    scripts = _REPO_ROOT / ".claude/skills/ppt-master/scripts"
    if str(scripts) not in sys.path:
        sys.path.insert(0, str(scripts))
    import gate_receipts
    import verify_deck

    checked: list[str] = []
    with tempfile.TemporaryDirectory(prefix="hash-receipts-") as scratch:
        project = Path(scratch)
        spec = project / "design_spec.md"
        lock = project / "spec_lock.md"
        spec.write_text("original spec", encoding="utf-8")
        lock.write_text("original lock", encoding="utf-8")
        stamp = project / ".spec_pass.json"
        stamp.write_text("not JSON", encoding="utf-8")
        assert not verify_deck._stamp_fresh(stamp, [spec, lock])
        gate_receipts.write_gate_pass(project, "spec", gate_receipts.capture_gate(project, "spec"))
        assert verify_deck._stamp_fresh(stamp, [spec, lock])
        checked.append("valid content-bound PASS is reusable and malformed PASS is rejected")

        original = spec.read_bytes()
        original_stat = spec.stat()
        spec.write_text("changed with identical mtime", encoding="utf-8")
        os.utime(spec, ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns))
        assert not verify_deck._stamp_fresh(stamp, [spec, lock])
        spec.write_bytes(original)
        lock_bytes = lock.read_bytes()
        lock.unlink()
        assert not verify_deck._stamp_fresh(stamp, [spec, lock])
        lock.write_bytes(lock_bytes)
        assert verify_deck._stamp_fresh(stamp, [spec, lock])
        checked.append("mtime-preserving edits and deleted required inputs invalidate PASS")

        with patch.object(gate_receipts, "_validator_bundle", return_value="0" * 64):
            assert not verify_deck._stamp_fresh(stamp, [spec, lock])
        assert not gate_receipts.gate_pass_current(project, "spec", options={"strict": True})
        before = gate_receipts.capture_gate(project, "spec")
        spec.write_text("changed while validator runs", encoding="utf-8")
        try:
            gate_receipts.write_gate_pass(project, "spec", before)
        except ValueError as exc:
            assert "changed during validation" in str(exc)
        else:
            raise AssertionError("changed inputs must not receive a PASS")
        checked.append("validator changes options and in-flight input edits invalidate PASS")

        svg_dir = project / "svg_output"
        svg_dir.mkdir()
        page = svg_dir / "P01.svg"
        page.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"/>', encoding="utf-8")
        gate_receipts.write_gate_pass(project, "svg-quality", gate_receipts.capture_gate(project, "svg-quality"))
        assert gate_receipts.gate_pass_current(project, "svg-quality")
        page.rename(svg_dir / "P02.svg")
        assert not gate_receipts.gate_pass_current(project, "svg-quality")
        checked.append("same page count with a changed filename invalidates SVG PASS")

        spec.write_text("## IX. Slide design\n#### Slide 01 — Cover\n- **Cover impact**: Clear opening\n", encoding="utf-8")
        lock.write_text(
            "## canvas\n- viewBox: 0 0 1280 720\n## mode\n- mode: briefing\n"
            "## visual_style\n- visual_style: minimal\n## colors\n"
            "- primary: #17324D\n- secondary: #306080\n- text: #172B4D\n"
            "- muted: #566778\n- background: #FFFFFF\n- border: #DDDDDD\n"
            "## typography\n- font_family: Pretendard\n- body: 24\n"
            "## page_rhythm\n- P01: anchor\n## pptx_structure\n- mode: flat\n",
            encoding="utf-8",
        )
        completed = subprocess.run(
            [sys.executable, str(scripts / "validate_spec.py"), str(project)],
            capture_output=True, text=True, check=False,
        )
        assert completed.returncode == 0, completed.stdout + completed.stderr
        assert gate_receipts.gate_pass_current(project, "spec")
        checked.append("real validate_spec producer writes a reusable hash-bound receipt")

        # Use the same valid flat-route planning contract for the actual full
        # SVG producer, so the receipt is earned through the owning checker.
        (svg_dir / "P02.svg").rename(page)
        page.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720" '
            'data-pptx-page-role="cover">'
            '<g id="content"><rect x="0" y="0" width="1280" height="720" fill="#FFFFFF"/>'
            '<text x="80" y="100" font-family="Pretendard" font-size="32" fill="#172B4D">'
            'Content receipt fixture</text></g></svg>', encoding="utf-8",
        )
        completed = subprocess.run(
            [sys.executable, str(scripts / "svg_quality_checker.py"), str(project)],
            capture_output=True, text=True, check=False,
        )
        assert completed.returncode == 0, completed.stdout + completed.stderr
        assert gate_receipts.gate_pass_current(project, "svg-quality")
        checked.append("real full SVG checker writes a reusable hash-bound receipt")
    return checked


def _verify_audit_gate_regressions(source: AgentPipeline) -> list[str]:
    """Exercise approval drift and rejection on isolated completed fixtures."""
    checked: list[str] = []
    with tempfile.TemporaryDirectory(prefix="audit-gates-") as scratch:
        root = Path(scratch)

        def clone(label: str) -> AgentPipeline:
            destination = root / label
            shutil.copytree(source.workspace, destination)
            return AgentPipeline(destination)

        def review_state(pipeline: AgentPipeline, target: str) -> None:
            state = pipeline.state
            index = max(i for i, entry in enumerate(state["history"]) if entry["to"] == target)
            state.update(status=target, revision=index + 1, history=state["history"][:index + 1])
            pipeline_module.write_json(pipeline.control_dir / "state.json", state)
            (pipeline.control_dir / "03_design/verification.json").unlink(missing_ok=True)

        for status in ("DESIGN_AUTHORIZED", "DESIGN_ACTIVE", "DESIGN_APPROVED", "DESIGN_REVIEW", "COMPLETE"):
            pipeline = clone("missing-approval-" + status)
            if status != "COMPLETE":
                review_state(pipeline, status)
            (pipeline.control_dir / "approvals/research_handoff_approval.json").unlink()
            findings = pipeline.validate()
            assert any("research handoff approval is required" in item for item in findings), findings
            if status == "DESIGN_APPROVED":
                _expect_blocked(pipeline.verify_design, "research handoff approval is required")
        checked.append("post-G2 missing approval blocks validation and verification")

        for status in ("DESIGN_HANDOFF_WAITING_USER", "DESIGN_AUTHORIZED"):
            pipeline = clone("upstream-" + status)
            review_state(pipeline, status)
            (pipeline.control_dir / "01_intent/approval.json").unlink()
            findings = pipeline.validate()
            assert any("intent approval is required" in item for item in findings), findings
        checked.append("handoff waiting and authorized states enforce upstream intent")

        pipeline = clone("wrong-bound-paths")
        approval_path = pipeline.control_dir / "approvals/research_handoff_approval.json"
        approval = json.loads(approval_path.read_text(encoding="utf-8"))
        bound = approval["bound_file_sha256"]
        original_path = next(iter(bound))
        replacement = pipeline.workspace / "unapproved-copy.md"
        replacement.write_bytes((pipeline.workspace / original_path).read_bytes())
        bound[replacement.name] = bound.pop(original_path)
        pipeline_module.write_json(approval_path, approval)
        state = pipeline.state
        state["research_handoff_approval_sha256"] = sha256_file(approval_path)
        pipeline_module.write_json(pipeline.control_dir / "state.json", state)
        _expect_blocked(pipeline._verify_research_handoff_approval, "bound file")
        checked.append("G2 binds the exact required path set")

        for relative in ("approvals/research_handoff_approval.json", "approvals/research_handoff_questions.json",
                         "approvals/research_handoff_answers.json"):
            pipeline = clone("malformed-" + Path(relative).stem)
            (pipeline.control_dir / relative).write_text("not JSON", encoding="utf-8")
            assert pipeline.validate(), relative
        checked.append("malformed G2 approval question and answer fail as findings")

        pipeline = clone("review-rejection")
        review_state(pipeline, "DESIGN_REVIEW")
        candidate = pipeline.control_dir / "03_design/verification_candidate.json"
        candidate_bytes = candidate.read_bytes()
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            rejected = pipeline.complete_design(contact_sheet_review="FAIL", review_note="Fix the crowded page before release")
        assert pipeline.state["status"] == "DESIGN_APPROVED"
        assert not candidate.exists()
        assert (pipeline.workspace / rejected["archive"] / candidate.name).read_bytes() == candidate_bytes

        def verify(command: list[str], **_: object) -> subprocess.CompletedProcess[str]:
            contact = pipeline.workspace / "_pptx_render/fixture-grid.png"
            contact.parent.mkdir(parents=True, exist_ok=True)
            contact.write_bytes(_valid_png())
            return subprocess.CompletedProcess(command, 0, stdout="PASS", stderr="")

        with patch("presentation_agents.pipeline._run_bounded_subprocess", side_effect=verify), patch(
            "presentation_agents.pipeline._officecli_bin", return_value="fixture-officecli"
        ):
            pipeline.verify_design()
        # A re-export can yield the same bytes and timestamp. Rejecting the
        # same candidate again must preserve both review decisions.
        candidate.write_bytes(candidate_bytes)
        state = pipeline.state
        state["verification_candidate_sha256"] = sha256_file(candidate)
        pipeline_module.write_json(pipeline.control_dir / "state.json", state)
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            repeated = pipeline.request_design_changes("The repeated candidate still needs correction")
        assert repeated["archive"] != rejected["archive"]
        assert (pipeline.workspace / rejected["archive"] / candidate.name).read_bytes() == candidate_bytes
        assert (pipeline.workspace / repeated["archive"] / candidate.name).read_bytes() == candidate_bytes
        with patch("presentation_agents.pipeline._run_bounded_subprocess", side_effect=verify), patch(
            "presentation_agents.pipeline._officecli_bin", return_value="fixture-officecli"
        ):
            pipeline.verify_design()
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            pipeline.complete_design(contact_sheet_review="PASS", review_note="Corrected candidate reviewed successfully")
        assert pipeline.validate() == []
        checked.append("repeated review FAIL preserves both candidates then verifies and completes normally")

        for kind in ("pptx", "spec", "missing-receipt"):
            pipeline = clone("stale-complete-" + kind)
            if kind == "pptx":
                (pipeline.workspace / "exports/fixture.pptx").write_bytes(b"stale candidate")
            elif kind == "spec":
                (pipeline.workspace / "design_spec.md").write_text("Changed design direction", encoding="utf-8")
            else:
                (pipeline.control_dir / "03_design/verification.json").unlink()
            with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
                pipeline.reopen_design("Repair the stale completed candidate")
            expected = "DESIGN_ACTIVE" if kind == "spec" else "DESIGN_APPROVED"
            assert pipeline.state["status"] == expected, pipeline.state
            assert not (pipeline.control_dir / "03_design/verification.json").exists()
        checked.append("stale or missing completion can reopen without a fabricated PASS")
    return checked


def main() -> int:
    if not __debug__:
        print("ERROR: validation must run without Python -O", file=sys.stderr)
        return 2
    os.environ.pop("PRESENTATION_SUITE_ACTOR", None)
    manifest_findings = validate_agent_manifests(_SUITE_ROOT / "agents")
    assert manifest_findings == [], manifest_findings
    slidemaster_dependencies = [
        _REPO_ROOT / ".claude/skills/ppt-master/SKILL.md",
        _REPO_ROOT / ".claude/skills/ppt-master/workflows/routing.md",
        _REPO_ROOT / ".claude/skills/ppt-master/scripts/validate_spec.py",
        _REPO_ROOT / ".claude/skills/ppt-master/scripts/svg_quality_checker.py",
        _REPO_ROOT / ".claude/skills/ppt-master/scripts/svg_to_pptx.py",
        _REPO_ROOT / ".claude/skills/ppt-master/scripts/verify_deck.py",
    ]
    missing_dependencies = [str(path) for path in slidemaster_dependencies if not path.is_file()]
    assert missing_dependencies == [], missing_dependencies
    swift = shutil.which("swift")
    quicklook_renderer = _SUITE_ROOT / "scripts/render_quicklook_contact_sheet.swift"
    if sys.platform == "darwin" and swift and quicklook_renderer.is_file():
        safe_geometry = subprocess.run(
            [swift, str(quicklook_renderer), "--validate-geometry", "10", "1920", "1080"],
            capture_output=True,
            text=True,
            check=False,
            timeout=120,
        )
        assert safe_geometry.returncode == 0, safe_geometry.stderr
        geometry = json.loads(safe_geometry.stdout)
        assert 1 <= geometry["columns"] <= 10
        assert geometry["canvas_pixels"] <= 40_000_000
        dangerous_geometry = subprocess.run(
            [swift, str(quicklook_renderer), "--validate-geometry", "500", "1", "10000"],
            capture_output=True,
            text=True,
            check=False,
            timeout=120,
        )
        assert dangerous_geometry.returncode == 1
        assert "unsupported slide geometry" in dangerous_geometry.stderr
    _expect_blocked(
        lambda: _validate_png_bytes(_illegal_truecolor_png()),
        "invalid contact-sheet PNG",
    )
    _expect_blocked(
        lambda: _validate_png_bytes(_indexed_png_without_palette()),
        "invalid contact-sheet PNG",
    )
    _expect_blocked(
        lambda: _validate_png_bytes(_nonconsecutive_idat_png()),
        "invalid contact-sheet PNG",
    )

    with tempfile.TemporaryDirectory(prefix="presentation-agent-suite-") as temp_dir:
        root = Path(temp_dir)
        _verify_bounded_subprocess_contract()
        _verify_officecli_render_failure_contract(root)
        _verify_quicklook_helper_contract(root)
        mutated_agents = root / "mutated-agents"
        shutil.copytree(_SUITE_ROOT / "agents", mutated_agents)
        mutated_manifest_path = mutated_agents / "01_intent_architect/agent.json"
        mutated_manifest = json.loads(mutated_manifest_path.read_text(encoding="utf-8"))
        mutated_manifest["order"] = "one"
        mutated_manifest["owned_outputs"] = ["arbitrary-output"]
        mutated_manifest_path.write_text(
            json.dumps(mutated_manifest, ensure_ascii=False),
            encoding="utf-8",
        )
        mutated_findings = validate_agent_manifests(mutated_agents)
        assert any("order must be an integer" in item for item in mutated_findings)
        assert any("owned_outputs do not match" in item for item in mutated_findings)

        pipeline = AgentPipeline.create(
            projects_root=root,
            slug="ivd-market-entry",
            request="5장 경영진 프레젠테이션을 준비한다.",
            canvas_format="ppt169",
            date_prefix="20260904",
        )
        assert pipeline.state["status"] == "INTENT_INTERVIEW"
        assert pipeline.current_agent_id == "intent-architect"

        incomplete = _sample_intent()
        incomplete_fields = incomplete["fields"]
        assert isinstance(incomplete_fields, dict)
        incomplete_fields.pop("desired_decision")
        _expect_blocked(
            lambda: pipeline.record_intent(incomplete, _sample_outline()),
            "missing intent fields",
        )

        pipeline.record_intent(_sample_intent(), _sample_outline())
        assert pipeline.state["status"] == "INTENT_REVIEW"
        blueprint = pipeline.workspace / "agent_pipeline/01_intent/presentation_blueprint.md"
        assert blueprint.is_file()
        assert "## 5. 슬라이드별 구조" in blueprint.read_text(encoding="utf-8")
        denied_approval = subprocess.run(
            [
                sys.executable,
                str(_AGENT_LAUNCHER),
                "approve-intent",
                str(pipeline.workspace),
                "--decision",
                "APPROVE",
                "--note",
                "agent must not approve",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert denied_approval.returncode == 1
        assert "supervisor-only" in denied_approval.stderr
        _expect_blocked(pipeline.start_research, "intent approval")
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            _expect_blocked(
                lambda: pipeline.approve_intent("approval without interview", "APPROVE"),
                "confirmation question bound to the current blueprint",
            )
        state_path = pipeline.workspace / "agent_pipeline/state.json"
        intent_questions_path = pipeline.workspace / "agent_pipeline/01_intent/questions.json"
        raise_question_before = {
            state_path: state_path.read_bytes(),
            intent_questions_path: intent_questions_path.read_bytes(),
        }
        real_write_json = pipeline_module.write_json
        raise_state_write_attempts = 0

        def _fail_raise_question_state_write(path: Path, value: object) -> None:
            nonlocal raise_state_write_attempts
            if path == state_path and raise_state_write_attempts == 0:
                raise_state_write_attempts += 1
                raise OSError("synthetic raise-question state write failure")
            real_write_json(path, value)

        with patch(
            "presentation_agents.pipeline.write_json",
            side_effect=_fail_raise_question_state_write,
        ):
            try:
                pipeline.raise_question(
                    stage="intent",
                    question="이 설계서의 목적·청중·5장 구조를 최종 승인합니까?",
                    impact="승인된 구조만 리서치 단계로 전달됩니다.",
                )
            except OSError as exc:
                assert "synthetic raise-question state write failure" in str(exc)
            else:
                raise AssertionError("raise-question state failure must abort the transaction")
        for path, before in raise_question_before.items():
            assert path.read_bytes() == before
        assert pipeline.validate() == []

        intent_question = pipeline.raise_question(
            stage="intent",
            question="이 설계서의 목적·청중·5장 구조를 최종 승인합니까?",
            impact="승인된 구조만 리서치 단계로 전달됩니다.",
        )
        _expect_blocked(
            lambda: pipeline.answer_question(
                str(intent_question["question_id"]), "agent-authored answer"
            ),
            "supervisor actor",
        )
        intent_questions_bytes = intent_questions_path.read_bytes()
        forged_questions = json.loads(intent_questions_bytes)
        forged_questions["questions"][0]["status"] = "answered"
        forged_questions["questions"][0]["answer"] = "agent-forged answer"
        intent_questions_path.write_text(json.dumps(forged_questions), encoding="utf-8")
        forged_findings = pipeline.validate()
        assert any("lacks a supervisor answer receipt" in item for item in forged_findings)
        intent_questions_path.write_bytes(intent_questions_bytes)
        intent_answer = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "answer-question",
                str(pipeline.workspace),
                "--question-id",
                str(intent_question["question_id"]),
                "--answer",
                "예, 전체 설계를 승인합니다.",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert intent_answer.returncode == 0, intent_answer.stderr
        assert pipeline.state["status"] == "INTENT_REVIEW"
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            try:
                pipeline.approve_intent("", "APPROVE")
            except ValueError as exc:
                assert "approval note" in str(exc)
            else:
                raise AssertionError("empty approval note must be rejected")

        outline_path = pipeline.workspace / "agent_pipeline/01_intent/story_outline.json"
        outline_bytes = outline_path.read_bytes()
        outline_path.write_text('{"slides": []}\n', encoding="utf-8")
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            _expect_blocked(
                lambda: pipeline.approve_intent("must reject tampered outline", "APPROVE"),
                "story outline slides are required",
            )
        outline_path.write_bytes(outline_bytes)

        supervisor_approval = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "approve-intent",
                str(pipeline.workspace),
                "--decision",
                "APPROVE",
                "--note",
                "approved by validation fixture",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_approval.returncode == 0, supervisor_approval.stderr
        approval = json.loads(supervisor_approval.stdout)
        assert approval["artifact_sha256"]
        request_path = pipeline.workspace / "agent_pipeline/01_intent/request.md"
        request_bytes = request_path.read_bytes()
        request_path.write_bytes(request_bytes + b"\nTAMPER\n")
        _expect_blocked(pipeline.start_research, "intent approval is stale")
        request_path.write_bytes(request_bytes)
        blueprint_bytes = blueprint.read_bytes()
        blueprint.write_bytes(blueprint_bytes + b"\nTAMPER\n")
        _expect_blocked(pipeline.start_research, "intent approval is stale")
        blueprint.write_bytes(blueprint_bytes)
        pipeline.start_research()
        assert pipeline.current_agent_id == "research-analyst"

        source_md = root / "regulator.md"
        source_md.write_text(
            "# 규제기관 공개 자료\n\n시장 A의 승인 절차는 세 단계이며 평균 처리 기간은 공개되지 않았다.\n",
            encoding="utf-8",
        )
        source_svg = root / "market-map.svg"
        source_svg.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="40" height="20">'
            '<rect width="40" height="20" fill="#1565C0"/></svg>\n',
            encoding="utf-8",
        )
        pipeline.add_text_source(
            source_id="SRC-001",
            title="규제기관 공개 자료",
            url="https://example.invalid/regulator",
            markdown_file=source_md,
            license_status="public-information",
            evidence_locator="#section=approval-process",
        )
        preserved_text = (
            pipeline.workspace / "agent_pipeline/02_research/text/SRC-001.md"
        )
        preserved_bytes = preserved_text.read_bytes()
        duplicate_md = root / "duplicate.md"
        duplicate_md.write_text("# 잘못된 덮어쓰기 후보\n", encoding="utf-8")
        _expect_blocked(
            lambda: pipeline.add_text_source(
                source_id="SRC-001",
                title="중복 ID",
                url="https://example.invalid/duplicate",
                markdown_file=duplicate_md,
                license_status="fixture-only",
                evidence_locator="#duplicate",
            ),
            "duplicate source_id",
        )
        assert preserved_text.read_bytes() == preserved_bytes
        source_manifest_path = (
            pipeline.workspace / "agent_pipeline/02_research/source_manifest.json"
        )
        manifest_before_invalid = source_manifest_path.read_bytes()
        try:
            pipeline.add_text_source(
                source_id="SRC-004",
                title="",
                url="https://example.invalid/invalid-metadata",
                markdown_file=duplicate_md,
                license_status="fixture-only",
                evidence_locator="#invalid",
            )
        except ValueError as exc:
            assert "title" in str(exc)
        else:
            raise AssertionError("invalid source metadata must be rejected")
        assert not (
            pipeline.workspace / "agent_pipeline/02_research/text/SRC-004.md"
        ).exists()
        assert source_manifest_path.read_bytes() == manifest_before_invalid

        transaction_destination = (
            pipeline.workspace / "agent_pipeline/02_research/text/SRC-004.md"
        )
        def _fail_source_manifest_write(path: Path, value: object) -> None:
            if path == source_manifest_path:
                raise OSError("synthetic manifest write failure")
            real_write_json(path, value)

        with patch(
            "presentation_agents.pipeline.write_json",
            side_effect=_fail_source_manifest_write,
        ):
            try:
                pipeline.add_text_source(
                    source_id="SRC-004",
                    title="트랜잭션 복구 자료",
                    url="https://example.invalid/transaction",
                    markdown_file=duplicate_md,
                    license_status="fixture-only",
                    evidence_locator="#transaction",
                )
            except OSError as exc:
                assert "synthetic manifest write failure" in str(exc)
            else:
                raise AssertionError("manifest failure must abort the source transaction")
        assert not transaction_destination.exists()
        assert source_manifest_path.read_bytes() == manifest_before_invalid
        assert list(transaction_destination.parent.glob(f".{transaction_destination.name}.*")) == []
        pipeline.add_text_source(
            source_id="SRC-004",
            title="트랜잭션 복구 자료",
            url="https://example.invalid/transaction",
            markdown_file=duplicate_md,
            license_status="fixture-only",
            evidence_locator="#transaction",
        )

        external_target = root / "outside.md"
        external_target.write_text("outside must remain unchanged", encoding="utf-8")
        symlink_destination = (
            pipeline.workspace / "agent_pipeline/02_research/text/SRC-005.md"
        )
        symlink_destination.symlink_to(external_target)
        _expect_blocked(
            lambda: pipeline.add_text_source(
                source_id="SRC-005",
                title="symlink probe",
                url="https://example.invalid/symlink",
                markdown_file=duplicate_md,
                license_status="fixture-only",
                evidence_locator="#symlink",
            ),
            "destination already exists",
        )
        assert external_target.read_text(encoding="utf-8") == "outside must remain unchanged"
        symlink_destination.unlink()

        pipeline.add_media_source(
            source_id="SRC-002",
            title="시장 지도",
            url="https://example.invalid/market-map.svg",
            media_file=source_svg,
            license_status="fixture-only",
            evidence_locator="#asset=market-map",
        )
        preserved_media = (
            pipeline.workspace / "agent_pipeline/02_research/media/SRC-002/market-map.svg"
        )
        preserved_media_bytes = preserved_media.read_bytes()
        _expect_blocked(
            lambda: pipeline.add_media_source(
                source_id="SRC-002",
                title="duplicate media",
                url="https://example.invalid/duplicate-media.svg",
                media_file=source_svg,
                license_status="fixture-only",
                evidence_locator="#duplicate-media",
            ),
            "duplicate source_id",
        )
        assert preserved_media.read_bytes() == preserved_media_bytes
        pipeline.add_blocked_source(
            source_id="SRC-003",
            title="접근 제한 경쟁사 보고서",
            url="https://example.invalid/restricted-report",
            resource_type="text",
            license_status="UNKNOWN",
            evidence_locator="#access=subscription-wall",
            failure_reason="구독 권한이 없어 본문을 내려받을 수 없음",
        )
        research_plan = pipeline.workspace / "agent_pipeline/02_research/research_plan.md"
        research_plan.write_text(
            research_plan.read_text(encoding="utf-8").replace("- [ ]", "- [x]"),
            encoding="utf-8",
        )

        question = pipeline.raise_question(
            stage="research",
            question="시장 매력도와 실행 가능성 중 어느 쪽을 더 우선할까요?",
            impact="가중치와 최종 순위가 바뀐다.",
        )
        analysis_md = root / "analysis.md"
        analysis_md.write_text(
            "# 시장 진입 분석\n\n"
            "## Executive conclusion\n\n시장 A를 1순위 검증 대상으로 본다.\n\n"
            "## 조사 질문별 답변\n\n후보 시장, 규모, 규제, 채널, 평가표와 실행 제약을 확인했다.\n\n"
            "## 핵심 주장과 근거\n\n승인 절차는 세 단계다. "
            "[SOURCE: SRC-001 @ #section=approval-process]\n\n"
            "## 비교·계산·가정\n\n실행 가능성 60%, 시장 매력도 40%를 합성 fixture의 가정으로 사용한다.\n\n"
            "## 반대 근거와 대안 해석\n\n평균 처리 기간이 비공개이므로 속도 우위는 확정하지 않는다.\n\n"
            "## 불확실성·누락·시간 경계\n\n구독 자료는 원문을 확보하지 못해 결론 근거에서 제외한다. "
            "[ACCESS_RESTRICTED: SRC-003]\n\n"
            "## 슬라이드별 evidence packet\n\n"
            "- [EVIDENCE: EVID-P01-01] 후보 시장 정의 [SOURCE: SRC-001 @ #section=approval-process]\n"
            "- [EVIDENCE: EVID-P02-01] 시장 규모 [SOURCE: SRC-001 @ #section=approval-process]\n"
            "- [EVIDENCE: EVID-P02-02] 전망 범위와 기준연도 [SOURCE: SRC-001 @ #section=approval-process]\n"
            "- [EVIDENCE: EVID-P03-01] 규제기관 원문 [SOURCE: SRC-001 @ #section=approval-process]\n"
            "- [EVIDENCE: EVID-P03-02] 채널 근거 [SOURCE: SRC-001 @ #section=approval-process]\n"
            "- [EVIDENCE: EVID-P04-01] 평가표 [SOURCE: SRC-001 @ #section=approval-process]\n"
            "- [EVIDENCE: EVID-P04-02] 가중치 근거 [SOURCE: SRC-001 @ #section=approval-process]\n"
            "- [EVIDENCE: EVID-P05-01] 실행 제약 [SOURCE: SRC-002 @ #asset=market-map]\n"
            "- [EVIDENCE: EVID-P05-02] 성공 기준 [SOURCE: SRC-002 @ #asset=market-map]\n\n"
            "## 출처 목록\n\nSRC-001과 SRC-004는 보존된 자료다. "
            "[SOURCE: SRC-004 @ #transaction] SRC-003은 접근 제한 기록이다.\n",
            encoding="utf-8",
        )
        guide_md = root / "design_guidelines.md"
        guide_md.write_text(
            "# 디자인 가이드라인\n\n"
            "## P01\n\n- Claim: 의사결정 질문\n- Evidence: [EVIDENCE: EVID-P01-01]\n- Visual: 단일 질문 프레임\n- Axes/units: 해당 없음\n- Media/license: 없음\n- Caveat/source note: 정의 범위 표시\n- Density/emphasis: 질문 우선\n- Avoid: 장식적 도형\n\n"
            "## P02\n\n- Claim: 시장 수요\n- Evidence: [EVIDENCE: EVID-P02-01] [EVIDENCE: EVID-P02-02]\n- Visual: small multiples\n- Axes/units: 기준연도와 단위 명시\n- Media/license: 없음\n- Caveat/source note: 전망 범위 표시\n- Density/emphasis: 수치 우선\n- Avoid: 다른 기준연도 혼합\n\n"
            "## P03\n\n- Claim: 규제와 채널\n- Evidence: [EVIDENCE: EVID-P03-01] [EVIDENCE: EVID-P03-02]\n- Visual: evidence matrix\n- Axes/units: 단계와 기간 분리\n- Media/license: 없음\n- Caveat/source note: [SOURCE: SRC-001 @ #section=approval-process]\n- Density/emphasis: 장벽 우선\n- Avoid: 미확인 기간 단정\n\n"
            "## P04\n\n- Claim: 시장 A 우선\n- Evidence: [EVIDENCE: EVID-P04-01] [EVIDENCE: EVID-P04-02]\n- Visual: ranked scorecard\n- Axes/units: 가중 점수\n- Media/license: 없음\n- Caveat/source note: 민감도 표시\n- Density/emphasis: 순위 우선\n- Avoid: 가정 은폐\n\n"
            "## P05\n\n- Claim: 90일 검증\n- Evidence: [EVIDENCE: EVID-P05-01] [EVIDENCE: EVID-P05-02]\n- Visual: roadmap\n- Axes/units: 일 단위\n- Media/license: [SOURCE: SRC-002 @ #asset=market-map]\n- Caveat/source note: 종료 조건 표시\n- Density/emphasis: 결정 gate 우선\n- Avoid: 확정 일정처럼 표현\n",
            encoding="utf-8",
        )
        _expect_blocked(
            lambda: pipeline.complete_research(analysis_md, guide_md),
            "open research questions",
        )

        research_questions_path = (
            pipeline.workspace / "agent_pipeline/02_research/questions.json"
        )
        research_answers_path = (
            pipeline.workspace / "agent_pipeline/02_research/user_answers.json"
        )
        state_path = pipeline.workspace / "agent_pipeline/state.json"
        question_transaction_before = {
            research_questions_path: research_questions_path.read_bytes(),
            research_answers_path: research_answers_path.read_bytes(),
            state_path: state_path.read_bytes(),
        }

        def _fail_answer_receipt_write(path: Path, value: object) -> None:
            if path == research_answers_path:
                raise OSError("synthetic answer receipt write failure")
            real_write_json(path, value)

        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}), patch(
            "presentation_agents.pipeline.write_json",
            side_effect=_fail_answer_receipt_write,
        ):
            try:
                pipeline.answer_question(
                    str(question["question_id"]),
                    "실행 가능성 60%, 시장 매력도 40%",
                )
            except OSError as exc:
                assert "synthetic answer receipt write failure" in str(exc)
            else:
                raise AssertionError("answer receipt failure must abort the answer transaction")
        for path, before in question_transaction_before.items():
            assert path.read_bytes() == before

        state_write_attempts = 0

        def _fail_first_state_write(path: Path, value: object) -> None:
            nonlocal state_write_attempts
            if path == state_path and state_write_attempts == 0:
                state_write_attempts += 1
                raise OSError("synthetic state transition write failure")
            real_write_json(path, value)

        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}), patch(
            "presentation_agents.pipeline.write_json",
            side_effect=_fail_first_state_write,
        ):
            try:
                pipeline.answer_question(
                    str(question["question_id"]),
                    "실행 가능성 60%, 시장 매력도 40%",
                )
            except OSError as exc:
                assert "synthetic state transition write failure" in str(exc)
            else:
                raise AssertionError("state failure must abort the answer transaction")
        for path, before in question_transaction_before.items():
            assert path.read_bytes() == before

        research_answer = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "answer-question",
                str(pipeline.workspace),
                "--question-id",
                str(question["question_id"]),
                "--answer",
                "실행 가능성 60%, 시장 매력도 40%",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert research_answer.returncode == 0, research_answer.stderr

        completed_plan = research_plan.read_text(encoding="utf-8")
        research_plan.write_text(
            completed_plan.replace("후보 시장 정의", "승인 범위와 무관한 대체 질문", 1),
            encoding="utf-8",
        )
        _expect_blocked(
            lambda: pipeline.complete_research(analysis_md, guide_md),
            "research plan does not exactly match the approved outline",
        )
        research_plan.write_text(completed_plan, encoding="utf-8")
        research_plan.write_text(completed_plan.replace("- [x]", "* [ ]", 1), encoding="utf-8")
        _expect_blocked(
            lambda: pipeline.complete_research(analysis_md, guide_md),
            "research plan has incomplete evidence items",
        )
        research_plan.write_text(completed_plan, encoding="utf-8")

        wrong_packet_locator = root / "wrong-packet-locator.md"
        analysis_text = analysis_md.read_text(encoding="utf-8")
        packet_start = analysis_text.index("## 슬라이드별 evidence packet")
        wrong_packet_locator.write_text(
            analysis_text[:packet_start]
            + analysis_text[packet_start:].replace(
                "[SOURCE: SRC-001 @ #section=approval-process]",
                "[SOURCE: SRC-001 @ #wrong-packet-locator]",
            ),
            encoding="utf-8",
        )
        _expect_blocked(
            lambda: pipeline.complete_research(wrong_packet_locator, guide_md),
            "evidence packet has an invalid source locator",
        )

        source_manifest_bytes = source_manifest_path.read_bytes()
        source_manifest = json.loads(source_manifest_bytes)
        source_manifest["sources"][0]["local_path"] = "README.md"
        source_manifest["sources"][0]["sha256"] = sha256_file(
            pipeline.workspace / "README.md"
        )
        source_manifest_path.write_text(
            json.dumps(source_manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        _expect_blocked(
            lambda: pipeline.complete_research(analysis_md, guide_md),
            "source path is not canonical",
        )
        source_manifest_path.write_bytes(source_manifest_bytes)

        empty_guideline = root / "empty-guideline.md"
        empty_guideline.write_text(
            guide_md.read_text(encoding="utf-8").replace(
                "- Claim: 의사결정 질문",
                "- Claim: ",
                1,
            ),
            encoding="utf-8",
        )
        _expect_blocked(
            lambda: pipeline.complete_research(analysis_md, empty_guideline),
            "has empty fields: Claim",
        )

        restricted_as_evidence = root / "restricted-as-evidence.md"
        restricted_as_evidence.write_text(
            analysis_md.read_text(encoding="utf-8").replace(
                "[ACCESS_RESTRICTED: SRC-003]",
                "[SOURCE: SRC-003 @ #access=subscription-wall]",
            ),
            encoding="utf-8",
        )
        _expect_blocked(
            lambda: pipeline.complete_research(restricted_as_evidence, guide_md),
            "access-restricted source cannot be cited as evidence",
        )
        wrong_heading_levels = root / "wrong-heading-levels.md"
        wrong_heading_levels.write_text(
            analysis_md.read_text(encoding="utf-8").replace("\n## ", "\n### "),
            encoding="utf-8",
        )
        _expect_blocked(
            lambda: pipeline.complete_research(wrong_heading_levels, guide_md),
            "analysis H2 section sequence",
        )
        receipt = pipeline.complete_research(analysis_md, guide_md)
        assert receipt["analysis_pdf_sha256"]
        analysis_pdf = pipeline.workspace / "agent_pipeline/02_research/analysis.pdf"
        assert analysis_pdf.is_file() and analysis_pdf.stat().st_size > 1000

        _expect_blocked(
            lambda: pipeline.reopen_research("independent review remediation"),
            "supervisor actor",
        )
        denied_reopen = subprocess.run(
            [
                sys.executable,
                str(_AGENT_LAUNCHER),
                "reopen-research",
                str(pipeline.workspace),
                "--note",
                "agent must not reopen a completed research package",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert denied_reopen.returncode == 1
        assert "supervisor-only" in denied_reopen.stderr
        supervisor_reopen = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "reopen-research",
                str(pipeline.workspace),
                "--note",
                "independent review found evidence gaps requiring remediation",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_reopen.returncode == 0, supervisor_reopen.stderr
        reopened_state = json.loads(supervisor_reopen.stdout)
        assert reopened_state["status"] == "RESEARCH_ACTIVE"
        assert pipeline.current_agent_id == "research-analyst"
        assert pipeline.state["history"][-1]["from"] == "DESIGN_READY"
        assert pipeline.state["history"][-1]["to"] == "RESEARCH_ACTIVE"
        receipt = pipeline.complete_research(analysis_md, guide_md)
        assert pipeline.state["status"] == "DESIGN_READY"

        try:
            import pymupdf as fitz
        except ImportError as exc:
            raise AssertionError("PyMuPDF is required for PDF validation") from exc
        with fitz.open(analysis_pdf) as document:
            assert document.page_count >= 1
            extracted = "\n".join(page.get_text() for page in document)
        assert "시장" in extracted

        glyph_markdown = root / "glyph-regression.md"
        glyph_markdown.write_text(
            "한글 55~79세 55.6% (D=L-P) D>0 D=0 D<0 `code` [brackets] + - * / : ; , . _ # · — “문장” → ①②③ •\n",
            encoding="utf-8",
        )
        glyph_pdf = root / "glyph-regression.pdf"
        render_markdown_pdf(glyph_markdown, glyph_pdf, _REPO_ROOT)
        with fitz.open(glyph_pdf) as document:
            spans = [
                span
                for block in document[0].get_text("dict")["blocks"]
                for line in block.get("lines", [])
                for span in line.get("spans", [])
            ]
        glyph_probe_symbols = set("%~()[]=<>+-*/`:_#.,;")
        symbol_fonts = {
            character: span["font"]
            for span in spans
            for character in span["text"]
            if character in glyph_probe_symbols
        }
        assert glyph_probe_symbols <= symbol_fonts.keys(), symbol_fonts
        assert all(
            symbol_fonts[character] == "Helvetica"
            for character in glyph_probe_symbols
        ), symbol_fonts
        unicode_probe_symbols = set("·—“”→①②③•")
        unicode_symbol_fonts = {
            character: span["font"]
            for span in spans
            for character in span["text"]
            if character in unicode_probe_symbols
        }
        assert unicode_probe_symbols <= unicode_symbol_fonts.keys(), unicode_symbol_fonts
        assert all(
            unicode_symbol_fonts[character] == "Dotum"
            for character in unicode_probe_symbols
        ), unicode_symbol_fonts

        long_markdown = root / "long-analysis.md"
        long_markdown.write_text(
            "# Pagination regression\n\n" + " ".join(f"TOKEN{index:04d}" for index in range(1000)),
            encoding="utf-8",
        )
        long_pdf = root / "long-analysis.pdf"
        render_markdown_pdf(long_markdown, long_pdf, _REPO_ROOT)
        with fitz.open(long_pdf) as document:
            long_text = "\n".join(page.get_text() for page in document)
        missing_tokens = [f"TOKEN{index:04d}" for index in range(1000) if f"TOKEN{index:04d}" not in long_text]
        assert missing_tokens == [], missing_tokens[:10]

        _expect_blocked(pipeline.prepare_design, "research handoff approval")
        _expect_blocked(
            lambda: pipeline.request_research_handoff_approval(
                "Approve the verified research package for design handoff?",
                "Design cannot start until the current research package is explicitly approved.",
            ),
            "supervisor actor",
        )
        denied_research_approval_request = subprocess.run(
            [
                sys.executable,
                str(_AGENT_LAUNCHER),
                "request-research-approval",
                str(pipeline.workspace),
                "--question",
                "Approve the verified research package for design handoff?",
                "--impact",
                "Design remains blocked without explicit approval.",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert denied_research_approval_request.returncode == 1
        assert "supervisor-only" in denied_research_approval_request.stderr
        research_receipt_bytes_before_approval = (
            pipeline.workspace / "agent_pipeline/02_research/research_receipt.json"
        ).read_bytes()
        supervisor_research_approval_request = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "request-research-approval",
                str(pipeline.workspace),
                "--question",
                "Approve the verified research package for design handoff?",
                "--impact",
                "Design remains blocked without explicit approval.",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_research_approval_request.returncode == 0, supervisor_research_approval_request.stderr
        research_approval_question = json.loads(supervisor_research_approval_request.stdout)
        assert research_approval_question["question_id"] == "Q-RESEARCH-HANDOFF-001"
        assert research_approval_question["question_type"] == "research-handoff-confirmation"
        assert research_approval_question["subject_sha256"] == receipt["artifact_sha256"]
        assert pipeline.state["status"] == "DESIGN_HANDOFF_WAITING_USER"
        supervisor_research_answer = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "answer-research-approval",
                str(pipeline.workspace),
                "--question-id",
                research_approval_question["question_id"],
                "--answer",
                "APPROVE",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_research_answer.returncode == 0, supervisor_research_answer.stderr
        answered_research_approval = json.loads(supervisor_research_answer.stdout)
        assert answered_research_approval["answer"] == "APPROVE"
        assert pipeline.state["status"] == "DESIGN_READY"
        assert (
            pipeline.workspace / "agent_pipeline/02_research/research_receipt.json"
        ).read_bytes() == research_receipt_bytes_before_approval
        assert pipeline._verify_research_receipt()["artifact_sha256"] == receipt["artifact_sha256"]
        _expect_blocked(pipeline.prepare_design, "research handoff approval")
        _expect_blocked(
            lambda: pipeline.approve_research_handoff("explicit fixture approval", "APPROVE"),
            "supervisor actor",
        )
        supervisor_research_approval = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "approve-research-handoff",
                str(pipeline.workspace),
                "--decision",
                "APPROVE",
                "--note",
                "explicit fixture approval",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_research_approval.returncode == 0, supervisor_research_approval.stderr
        research_handoff_approval = json.loads(supervisor_research_approval.stdout)
        assert research_handoff_approval["status"] == "APPROVED"
        assert research_handoff_approval["decision"] == "APPROVE"
        assert research_handoff_approval["actor"] == "supervisor"
        assert research_handoff_approval["research_artifact_sha256"] == receipt["artifact_sha256"]
        assert research_handoff_approval["combined_artifact_sha256"] == receipt["artifact_sha256"]
        assert len(research_handoff_approval["bound_file_sha256"]) == 8
        assert pipeline.state["status"] == "DESIGN_AUTHORIZED"

        pdf_bytes = analysis_pdf.read_bytes()
        analysis_pdf.write_bytes(pdf_bytes + b"TAMPER")
        _expect_blocked(pipeline.prepare_design, "research receipt is stale")
        analysis_pdf.write_bytes(pdf_bytes)
        handoff = pipeline.prepare_design()
        assert pipeline.current_agent_id == "presentation-designer"
        assert pipeline.state["status"] == "DESIGN_ACTIVE"
        handoff_text = handoff.read_text(encoding="utf-8")
        assert ".claude/skills/ppt-master/workflows/routing.md" in handoff_text
        assert approval["artifact_sha256"] in handoff_text
        assert receipt["analysis_pdf_sha256"] in handoff_text
        _expect_blocked(pipeline.verify_design, "design direction approval")

        design_spec_path = pipeline.workspace / "design_spec.md"
        design_spec_path.write_text(
            "# Design Spec\n\nSynthetic contract validation fixture.\n",
            encoding="utf-8",
        )
        (pipeline.workspace / "spec_lock.md").write_text(
            "# Spec Lock\n\n- route: main-svg-generation\n",
            encoding="utf-8",
        )
        _expect_blocked(
            lambda: pipeline.approve_design("agent-authored approval", "APPROVE"),
            "supervisor actor",
        )
        denied_design_approval = subprocess.run(
            [
                sys.executable,
                str(_AGENT_LAUNCHER),
                "approve-design",
                str(pipeline.workspace),
                "--decision",
                "APPROVE",
                "--note",
                "agent must not approve",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert denied_design_approval.returncode == 1
        assert "supervisor-only" in denied_design_approval.stderr
        supervisor_design_approval = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "approve-design",
                str(pipeline.workspace),
                "--decision",
                "APPROVE",
                "--note",
                "synthetic user approval of the design direction",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_design_approval.returncode == 0, supervisor_design_approval.stderr
        assert pipeline.state["status"] == "DESIGN_APPROVED"
        design_mirror = pipeline.workspace / "sources/02_research_analysis.md"
        design_mirror_bytes = design_mirror.read_bytes()
        design_mirror.write_bytes(design_mirror_bytes + b"\nTAMPER\n")
        _expect_blocked(pipeline.verify_design, "design_handoff_sha256")
        design_mirror.write_bytes(design_mirror_bytes)
        design_spec_bytes = design_spec_path.read_bytes()
        design_spec_path.write_bytes(design_spec_bytes + b"TAMPER")
        _expect_blocked(pipeline.verify_design, "design direction approval is stale")
        design_spec_path.write_bytes(design_spec_bytes)
        svg_path = pipeline.workspace / "svg_output/01_cover.svg"
        svg_payload = (
            '<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" '
            'viewBox="0 0 1920 1080"><rect width="1920" height="1080" fill="#fff"/>'
            '<text x="120" y="180">Agent contract fixture</text></svg>'
        )
        svg_path.write_text(svg_payload, encoding="utf-8")
        pptx_path = pipeline.workspace / "exports/fixture.pptx"
        with zipfile.ZipFile(pptx_path, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("ppt/presentation.xml", "<p:presentation xmlns:p='p'/>")
        svg_intermediate = pipeline.workspace / "exports/fixture_svg.pptx"
        svg_intermediate.write_bytes(pptx_path.read_bytes())
        root_shadow = pipeline.workspace / "shadow.pptx"
        root_shadow.write_bytes(pptx_path.read_bytes())
        future_ns = pptx_path.stat().st_mtime_ns + 1_000_000_000
        os.utime(root_shadow, ns=(future_ns, future_ns))
        _expect_blocked(pipeline.verify_design, "root-level PPTX shadows")
        root_shadow.unlink()

        equal_time_export = pipeline.workspace / "exports/equal-time.pptx"
        equal_time_export.write_bytes(pptx_path.read_bytes())
        tied_ns = max(pptx_path.stat().st_mtime_ns, equal_time_export.stat().st_mtime_ns)
        os.utime(pptx_path, ns=(tied_ns, tied_ns))
        os.utime(equal_time_export, ns=(tied_ns, tied_ns))
        _expect_blocked(pipeline.verify_design, "ambiguous newest native exports")
        equal_time_export.unlink()

        compressed_bomb = pipeline.workspace / "exports/compressed-bomb.pptx"
        with zipfile.ZipFile(compressed_bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("ppt/presentation.xml", "<p:presentation xmlns:p='p'/>")
            archive.writestr("ppt/media/bomb.bin", b"0" * 1_000_000)
        future_ns = pptx_path.stat().st_mtime_ns + 1_000_000_000
        os.utime(compressed_bomb, ns=(future_ns, future_ns))
        _expect_blocked(pipeline.verify_design, "suspicious compression ratio")
        compressed_bomb.unlink()

        contact_sheet = pipeline.workspace / "_pptx_render/fixture-grid.png"
        contact_sheet.parent.mkdir(parents=True, exist_ok=True)
        contact_sheet.write_bytes(b"stale-grid")

        def _mock_verify_without_render(
            command: list[str], **_: object
        ) -> subprocess.CompletedProcess[str]:
            return subprocess.CompletedProcess(command, 0, stdout="PASS", stderr="")

        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_mock_verify_without_render,
        ), patch(
            "presentation_agents.pipeline._officecli_bin",
            return_value="/opt/homebrew/bin/officecli",
            create=True,
        ), patch(
            "presentation_agents.pipeline._render_macos_quicklook_contact_sheet",
            return_value=True,
        ) as forbidden_fallback:
            _expect_blocked(
                pipeline.verify_design,
                "OfficeCLI was available but did not render",
            )
        assert forbidden_fallback.call_count == 0

        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_mock_verify_without_render,
        ), patch(
            "presentation_agents.pipeline._officecli_bin",
            return_value=None,
            create=True,
        ), patch(
            "presentation_agents.pipeline._render_macos_quicklook_contact_sheet",
            return_value=False,
            create=True,
        ):
            _expect_blocked(
                pipeline.verify_design,
                "passed but did not render the selected PPTX",
            )

        def _mock_quicklook_fallback(
            pptx: Path,
            output: Path,
            *,
            diagnostics: list[str] | None = None,
        ) -> bool:
            assert pptx == pptx_path
            assert output == contact_sheet
            assert diagnostics == []
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_bytes(_valid_png())
            return True

        original_fallback_pptx_bytes = pptx_path.read_bytes()

        def _mutating_quicklook_fallback(
            pptx: Path,
            output: Path,
            *,
            diagnostics: list[str] | None = None,
        ) -> bool:
            success = _mock_quicklook_fallback(
                pptx,
                output,
                diagnostics=diagnostics,
            )
            with zipfile.ZipFile(pptx_path, "a", compression=zipfile.ZIP_STORED) as archive:
                archive.writestr("docProps/fallback-replacement-probe.xml", "<replacement/>")
            return success

        try:
            with patch(
                "presentation_agents.pipeline._run_bounded_subprocess",
                side_effect=_mock_verify_without_render,
            ), patch(
                "presentation_agents.pipeline._officecli_bin",
                return_value=None,
            ), patch(
                "presentation_agents.pipeline._render_macos_quicklook_contact_sheet",
                side_effect=_mutating_quicklook_fallback,
            ):
                _expect_blocked(
                    pipeline.verify_design,
                    "native PPTX changed during verification",
                )
        finally:
            pptx_path.write_bytes(original_fallback_pptx_bytes)
            contact_sheet.unlink(missing_ok=True)

        design_approved_state_bytes = state_path.read_bytes()
        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_mock_verify_without_render,
        ), patch(
            "presentation_agents.pipeline._officecli_bin",
            return_value=None,
        ), patch(
            "presentation_agents.pipeline._render_macos_quicklook_contact_sheet",
            side_effect=_mock_quicklook_fallback,
            create=True,
        ) as quicklook_fallback:
            fallback_candidate = pipeline.verify_design()
        assert quicklook_fallback.call_count == 1
        assert fallback_candidate["status"] == "AWAITING_CONTACT_SHEET_REVIEW"
        assert fallback_candidate["render_backend"] == "macos-quicklook-webkit"
        assert pipeline.state["status"] == "DESIGN_REVIEW"

        # Restore the pre-review state so the remaining failure-contract probes
        # continue to exercise verify_design independently.
        state_path.write_bytes(design_approved_state_bytes)
        candidate_path = pipeline.workspace / "agent_pipeline/03_design/verification_candidate.json"
        candidate_path.unlink()
        contact_sheet.unlink()
        pipeline = AgentPipeline(pipeline.workspace)

        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=subprocess.TimeoutExpired(["verify_deck.py"], 600),
        ):
            _expect_blocked(pipeline.verify_design, "timed out after 600 seconds")

        def _mock_verify_invalid_png(
            command: list[str], **_: object
        ) -> subprocess.CompletedProcess[str]:
            contact_sheet.parent.mkdir(parents=True, exist_ok=True)
            contact_sheet.write_bytes(
                b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
                b"\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
            )
            return subprocess.CompletedProcess(command, 0, stdout="PASS", stderr="")

        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_mock_verify_invalid_png,
        ), patch(
            "presentation_agents.pipeline._officecli_bin",
            return_value="/opt/homebrew/bin/officecli",
        ):
            _expect_blocked(pipeline.verify_design, "invalid contact-sheet PNG")

        def _mock_verify_oversized_png(
            command: list[str], **_: object
        ) -> subprocess.CompletedProcess[str]:
            contact_sheet.parent.mkdir(parents=True, exist_ok=True)
            with contact_sheet.open("wb") as handle:
                handle.truncate(250_000_001)
            return subprocess.CompletedProcess(command, 0, stdout="PASS", stderr="")

        real_path_read_bytes = Path.read_bytes

        def _reject_oversized_read(path: Path) -> bytes:
            if path == contact_sheet:
                raise AssertionError("oversized contact sheet must be rejected before read_bytes")
            return real_path_read_bytes(path)

        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_mock_verify_oversized_png,
        ), patch(
            "presentation_agents.pipeline._officecli_bin",
            return_value="/opt/homebrew/bin/officecli",
        ), patch.object(Path, "read_bytes", new=_reject_oversized_read):
            _expect_blocked(pipeline.verify_design, "oversized contact-sheet PNG")
        contact_sheet.unlink(missing_ok=True)

        def _mock_verify(command: list[str], **_: object) -> subprocess.CompletedProcess[str]:
            assert command[-1] == "--require-render-success"
            contact_sheet.parent.mkdir(parents=True, exist_ok=True)
            contact_sheet.write_bytes(_valid_png())
            return subprocess.CompletedProcess(command, 0, stdout="PASS", stderr="")

        original_pptx_bytes = pptx_path.read_bytes()

        def _mock_verify_after_pptx_replacement(
            command: list[str], **_: object
        ) -> subprocess.CompletedProcess[str]:
            result = _mock_verify(command)
            with zipfile.ZipFile(pptx_path, "a", compression=zipfile.ZIP_STORED) as archive:
                archive.writestr("docProps/replacement-probe.xml", "<replacement/>")
            return result

        try:
            with patch(
                "presentation_agents.pipeline._run_bounded_subprocess",
                side_effect=_mock_verify_after_pptx_replacement,
            ), patch(
                "presentation_agents.pipeline._officecli_bin",
                return_value="/opt/homebrew/bin/officecli",
            ):
                _expect_blocked(
                    pipeline.verify_design,
                    "native PPTX changed during verification",
                )
        finally:
            pptx_path.write_bytes(original_pptx_bytes)
            contact_sheet.unlink(missing_ok=True)

        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_mock_verify,
        ), patch(
            "presentation_agents.pipeline._officecli_bin",
            return_value=None,
        ):
            _expect_blocked(
                pipeline.verify_design,
                "contact sheet without OfficeCLI",
            )

        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_mock_verify,
        ), patch(
            "presentation_agents.pipeline._officecli_bin",
            return_value="/opt/homebrew/bin/officecli",
        ):
            candidate = pipeline.verify_design()
        assert candidate["status"] == "AWAITING_CONTACT_SHEET_REVIEW"
        assert candidate["render_backend"] == "officecli"
        assert candidate["pptx_path"] == "exports/fixture.pptx"
        assert candidate["design_approval_sha256"]
        assert pipeline.state["status"] == "DESIGN_REVIEW"

        candidate_path = (
            pipeline.workspace / "agent_pipeline/03_design/verification_candidate.json"
        )
        candidate_bytes = candidate_path.read_bytes()
        retargeted_candidate = json.loads(candidate_bytes)
        retargeted_candidate.update(
            {
                "route": "tampered-route",
                "owner": "tampered-owner",
                "svg_count": 999,
                "pptx_path": "README.md",
                "pptx_sha256": sha256_file(pipeline.workspace / "README.md"),
            }
        )
        candidate_path.write_text(
            json.dumps(retargeted_candidate, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            _expect_blocked(
                lambda: pipeline.complete_design(
                    contact_sheet_review="PASS",
                    review_note="synthetic candidate retarget review",
                ),
                "design verification candidate hash is stale",
            )
        candidate_path.write_bytes(candidate_bytes)

        _expect_blocked(
            lambda: pipeline.complete_design(
                contact_sheet_review="PASS",
                review_note="agent-authored review",
            ),
            "supervisor actor",
        )
        svg_path.write_text(svg_payload + "\n<!-- tampered -->\n", encoding="utf-8")
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            _expect_blocked(
                lambda: pipeline.complete_design(
                    contact_sheet_review="PASS",
                    review_note="synthetic contact-sheet review",
                ),
                "svg_bundle_sha256",
            )
        svg_path.write_text(svg_payload, encoding="utf-8")
        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}):
            _expect_blocked(
                lambda: pipeline.complete_design(
                    contact_sheet_review="PASS",
                    review_note="short",
                ),
                "at least 10 characters",
            )
        denied_completion = subprocess.run(
            [
                sys.executable,
                str(_AGENT_LAUNCHER),
                "complete-design",
                str(pipeline.workspace),
                "--contact-sheet-review",
                "PASS",
                "--review-note",
                "agent must not approve completion",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert denied_completion.returncode == 1
        assert "supervisor-only" in denied_completion.stderr
        final_receipt_path = pipeline.workspace / "agent_pipeline/03_design/verification.json"
        completion_before = state_path.read_bytes()
        completion_state_write_attempts = 0

        def _fail_complete_design_state_write(path: Path, value: object) -> None:
            nonlocal completion_state_write_attempts
            if path == state_path and completion_state_write_attempts == 0:
                completion_state_write_attempts += 1
                raise OSError("synthetic complete-design state write failure")
            real_write_json(path, value)

        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}), patch(
            "presentation_agents.pipeline.write_json",
            side_effect=_fail_complete_design_state_write,
        ):
            try:
                pipeline.complete_design(
                    contact_sheet_review="PASS",
                    review_note="synthetic rollback verification review",
                )
            except OSError as exc:
                assert "synthetic complete-design state write failure" in str(exc)
            else:
                raise AssertionError("complete-design state failure must abort the transaction")
        assert state_path.read_bytes() == completion_before
        assert not final_receipt_path.exists()
        assert pipeline.validate() == []

        forged_early_receipt = {
            **candidate,
            "status": "COMPLETE",
            "completed_at": "fixture",
            "verification_candidate_sha256": pipeline.state["verification_candidate_sha256"],
            "contact_sheet_review": "PASS",
            "contact_sheet_review_note": "synthetic early receipt probe",
        }
        real_write_json(final_receipt_path, forged_early_receipt)
        early_receipt_findings = pipeline.validate()
        assert any(
            "design verification receipt exists outside COMPLETE state" in finding
            for finding in early_receipt_findings
        ), early_receipt_findings
        final_receipt_path.unlink()
        assert pipeline.validate() == []

        supervisor_completion = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "complete-design",
                str(pipeline.workspace),
                "--contact-sheet-review",
                "PASS",
                "--review-note",
                "synthetic contact-sheet review",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_completion.returncode == 0, supervisor_completion.stderr
        final_receipt = json.loads(supervisor_completion.stdout)
        assert final_receipt["status"] == "COMPLETE"

        findings = pipeline.validate()
        assert findings == [], findings
        final_receipt_bytes = final_receipt_path.read_bytes()
        tampered_receipt = json.loads(final_receipt_bytes)
        tampered_receipt["route"] = "tampered-route"
        final_receipt_path.write_text(json.dumps(tampered_receipt), encoding="utf-8")
        receipt_findings = pipeline.validate()
        assert any("stale: route" in finding for finding in receipt_findings), receipt_findings
        final_receipt_path.write_bytes(final_receipt_bytes)
        assert pipeline.validate() == []
        state_path = pipeline.workspace / "agent_pipeline/state.json"
        state_bytes = state_path.read_bytes()
        tampered_state = json.loads(state_bytes)
        tampered_state["status"] = "INTENT_INTERVIEW"
        state_path.write_text(json.dumps(tampered_state), encoding="utf-8")
        rollback_findings = pipeline.validate()
        assert any("history tail" in finding for finding in rollback_findings), rollback_findings
        state_path.write_bytes(state_bytes)
        assert pipeline.validate() == []

        legacy_candidate = json.loads(candidate_path.read_text(encoding="utf-8"))
        legacy_candidate.pop("render_backend")
        candidate_path.write_text(
            json.dumps(legacy_candidate, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        legacy_candidate_sha = sha256_file(candidate_path)
        legacy_receipt = json.loads(final_receipt_path.read_text(encoding="utf-8"))
        legacy_receipt.pop("render_backend")
        legacy_receipt["verification_candidate_sha256"] = legacy_candidate_sha
        final_receipt_path.write_text(
            json.dumps(legacy_receipt, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        legacy_state = json.loads(state_path.read_text(encoding="utf-8"))
        legacy_state["verification_candidate_sha256"] = legacy_candidate_sha
        state_path.write_text(
            json.dumps(legacy_state, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        pipeline = AgentPipeline(pipeline.workspace)
        assert any("invalid render_backend" in item for item in pipeline.validate())

        _expect_blocked(
            lambda: pipeline.reopen_design("independent code review remediation"),
            "supervisor actor",
        )
        denied_design_reopen = subprocess.run(
            [
                sys.executable,
                str(_AGENT_LAUNCHER),
                "reopen-design",
                str(pipeline.workspace),
                "--note",
                "agent must not reopen a completed design",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert denied_design_reopen.returncode == 1
        prior_candidate_bytes = candidate_path.read_bytes()
        prior_receipt_bytes = final_receipt_path.read_bytes()
        prior_receipt_sha = hashlib.sha256(prior_receipt_bytes).hexdigest()
        first_history_root = (
            pipeline.workspace / "agent_pipeline/03_design/verification_history"
        )
        first_archive = first_history_root / prior_receipt_sha
        assert not first_history_root.exists()
        first_reopen_state_bytes = state_path.read_bytes()
        real_path_mkdir = Path.mkdir
        real_path_write_bytes = Path.write_bytes
        real_path_unlink = Path.unlink

        for failure_point in (
            "archive-mkdir",
            "archive-candidate-write",
            "archive-receipt-write",
            "reopen-record-write",
            "canonical-candidate-unlink",
            "canonical-receipt-unlink",
        ):
            injected = False

            def _fail_mkdir(
                path: Path,
                mode: int = 0o777,
                parents: bool = False,
                exist_ok: bool = False,
            ) -> None:
                nonlocal injected
                if failure_point == "archive-mkdir" and path == first_archive and not injected:
                    injected = True
                    raise OSError(f"synthetic {failure_point} failure")
                real_path_mkdir(path, mode=mode, parents=parents, exist_ok=exist_ok)

            def _fail_write_bytes(path: Path, payload: bytes) -> int:
                nonlocal injected
                target_name = {
                    "archive-candidate-write": "verification_candidate.json",
                    "archive-receipt-write": "verification.json",
                }.get(failure_point)
                if path.parent == first_archive and path.name == target_name and not injected:
                    injected = True
                    raise OSError(f"synthetic {failure_point} failure")
                return real_path_write_bytes(path, payload)

            def _fail_write_json(path: Path, value: object) -> None:
                nonlocal injected
                if (
                    failure_point == "reopen-record-write"
                    and path == first_archive / "reopen.json"
                    and not injected
                ):
                    injected = True
                    raise OSError(f"synthetic {failure_point} failure")
                real_write_json(path, value)

            def _fail_unlink(path: Path, missing_ok: bool = False) -> None:
                nonlocal injected
                target = {
                    "canonical-candidate-unlink": candidate_path,
                    "canonical-receipt-unlink": final_receipt_path,
                }.get(failure_point)
                if path == target and not injected:
                    injected = True
                    raise OSError(f"synthetic {failure_point} failure")
                real_path_unlink(path, missing_ok=missing_ok)

            with patch.dict(
                os.environ,
                {"PRESENTATION_SUITE_ACTOR": "supervisor"},
            ), patch.object(Path, "mkdir", new=_fail_mkdir), patch.object(
                Path,
                "write_bytes",
                new=_fail_write_bytes,
            ), patch.object(
                Path,
                "unlink",
                new=_fail_unlink,
            ), patch(
                "presentation_agents.pipeline.write_json",
                side_effect=_fail_write_json,
            ):
                try:
                    pipeline.reopen_design(f"synthetic {failure_point} remediation")
                except OSError as exc:
                    assert f"synthetic {failure_point} failure" in str(exc)
                else:
                    raise AssertionError(f"{failure_point} must abort reopen-design")
            assert injected
            assert state_path.read_bytes() == first_reopen_state_bytes
            assert candidate_path.read_bytes() == prior_candidate_bytes
            assert final_receipt_path.read_bytes() == prior_receipt_bytes
            assert not first_history_root.exists()

        supervisor_design_reopen = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "reopen-design",
                str(pipeline.workspace),
                "--note",
                "independent code review remediation",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_design_reopen.returncode == 0, supervisor_design_reopen.stderr
        pipeline = AgentPipeline(pipeline.workspace)
        assert pipeline.state["status"] == "DESIGN_APPROVED"
        assert not candidate_path.exists()
        assert not final_receipt_path.exists()
        archive = (
            pipeline.workspace
            / "agent_pipeline/03_design/verification_history"
            / prior_receipt_sha
        )
        assert (archive / "verification_candidate.json").read_bytes() == prior_candidate_bytes
        assert (archive / "verification.json").read_bytes() == prior_receipt_bytes
        reopen_record = json.loads((archive / "reopen.json").read_text(encoding="utf-8"))
        assert reopen_record["note"] == "independent code review remediation"
        assert reopen_record["prior_verification_sha256"] == prior_receipt_sha
        assert pipeline.validate() == []

        with patch(
            "presentation_agents.pipeline._run_bounded_subprocess",
            side_effect=_mock_verify,
        ), patch(
            "presentation_agents.pipeline._officecli_bin",
            return_value="/opt/homebrew/bin/officecli",
        ):
            reopened_candidate = pipeline.verify_design()
        assert reopened_candidate["render_backend"] == "officecli"
        supervisor_recompletion = subprocess.run(
            [
                sys.executable,
                str(_SUPERVISOR_LAUNCHER),
                "complete-design",
                str(pipeline.workspace),
                "--contact-sheet-review",
                "PASS",
                "--review-note",
                "synthetic post-remediation contact-sheet review",
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        assert supervisor_recompletion.returncode == 0, supervisor_recompletion.stderr
        assert pipeline.state["status"] == "COMPLETE"
        findings = pipeline.validate()
        assert findings == [], findings

        rollback_state_bytes = state_path.read_bytes()
        rollback_candidate_bytes = candidate_path.read_bytes()
        rollback_receipt_bytes = final_receipt_path.read_bytes()
        rollback_receipt_sha = hashlib.sha256(rollback_receipt_bytes).hexdigest()
        rollback_archive = (
            pipeline.workspace
            / "agent_pipeline/03_design/verification_history"
            / rollback_receipt_sha
        )
        reopen_state_write_attempts = 0

        def _fail_reopen_state_write(path: Path, value: object) -> None:
            nonlocal reopen_state_write_attempts
            if path == state_path and reopen_state_write_attempts == 0:
                reopen_state_write_attempts += 1
                raise OSError("synthetic reopen-design state write failure")
            real_write_json(path, value)

        with patch.dict(os.environ, {"PRESENTATION_SUITE_ACTOR": "supervisor"}), patch(
            "presentation_agents.pipeline.write_json",
            side_effect=_fail_reopen_state_write,
        ):
            try:
                pipeline.reopen_design("synthetic rollback remediation review")
            except OSError as exc:
                assert "synthetic reopen-design state write failure" in str(exc)
            else:
                raise AssertionError("reopen-design state failure must abort the transaction")
        assert state_path.read_bytes() == rollback_state_bytes
        assert candidate_path.read_bytes() == rollback_candidate_bytes
        assert final_receipt_path.read_bytes() == rollback_receipt_bytes
        assert not rollback_archive.exists()
        assert pipeline.validate() == []
        audit_checks = _verify_audit_gate_regressions(pipeline)
        cache_checks = _verify_hash_receipt_regressions()
        summary = {
            "agents": 3,
            "manifest_findings": len(manifest_findings),
            "slidemaster_dependencies": len(slidemaster_dependencies),
            "workspace": pipeline.workspace.name,
            "state": pipeline.state["status"],
            "sources": 4,
            "analysis_pdf_pages": receipt["analysis_pdf_pages"],
            "design_verification_contract": "mocked external verifier; real hashes",
            "validation_findings": len(findings),
            "audit_gate_regressions": audit_checks,
            "hash_cache_regressions": cache_checks,
        }
        print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
